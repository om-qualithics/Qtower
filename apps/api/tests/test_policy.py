import io
import uuid
from datetime import datetime, timezone

import docx
import pytest

from apps.api.core import storage
from apps.api.core.db import SessionLocal, org_scoped_session
from apps.api.modules.identity.models import Org, User
from apps.api.modules.policy import docgen, service
from apps.api.modules.policy.constants import MAX_UPLOAD_SIZE_BYTES
from apps.api.modules.policy.models import Policy
from apps.api.modules.policy.questions import default_answers
from apps.api.modules.policy.service import PolicyApprovalError, PolicyUploadError, PolicyValidationError


def _make_org(name: str) -> Org:
    db = SessionLocal()
    try:
        org = Org(name=name)
        db.add(org)
        db.commit()
        db.refresh(org)
        return org
    finally:
        db.close()


def _make_user(org: Org, email: str) -> User:
    with org_scoped_session(str(org.id)) as db:
        user = User(org_id=org.id, email=email)
        db.add(user)
        db.flush()
        db.refresh(user)
        db.expunge(user)
        return user


def _make_policy_with_document(org: Org, *, status: str = "draft") -> Policy:
    """A policy with a document already attached (storage_key set) - a
    real minimal docx is uploaded to that key (not just a fake string),
    since approve() now downloads and extracts it to cache a Markdown
    copy (Milestone 10)."""
    with org_scoped_session(str(org.id)) as db:
        policy = Policy(org_id=org.id, answers={}, status=status, storage_key="")
        db.add(policy)
        db.flush()
        db.refresh(policy)
        policy_id = str(policy.id)
        storage_key = f"policies/{org.id}/{policy_id}/document.docx"
        policy.storage_key = storage_key
        db.flush()
        db.refresh(policy)
        db.expunge(policy)

    doc = docx.Document()
    doc.add_paragraph("Test policy body text.")
    buf = io.BytesIO()
    doc.save(buf)
    storage.upload_bytes(storage_key, buf.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")

    return policy


def _delete_org(org: Org) -> None:
    with org_scoped_session(str(org.id)) as db:
        db.query(Policy).filter(Policy.org_id == org.id).delete(synchronize_session=False)
        db.query(User).filter(User.org_id == org.id).delete(synchronize_session=False)
    db = SessionLocal()
    try:
        db.query(Org).filter(Org.id == org.id).delete(synchronize_session=False)
        db.commit()
    finally:
        db.close()


def test_render_places_answers_at_the_right_insertion_points() -> None:
    """Hits the real MinIO-seeded template (same pattern as test_rls.py
    hitting a real Postgres) since this is the actual proof that the
    ~20 insertion points land correctly - a mock template would only
    prove the code runs, not that it's correct."""
    org = Org(id=uuid.uuid4(), name="Render Test Org")
    answers = default_answers()
    answers["q2_1_governance_owner"]["rows"] = [{"name": "Jane Doe", "title": "Chief AI Officer"}]
    answers["q2_4_approvers"]["rows"] = [{"name": "John Smith", "title": "CEO"}]
    answers["q3_9_default_tier"]["selected"] = "restricted"
    answers["q1_2_ai_systems"]["other"].append("Custom internal AI copilot")

    policy = Policy(
        id=uuid.uuid4(),
        org_id=org.id,
        status="draft",
        current_step=7,
        policy_owner_name="Jane Doe",
        approver_name="John Smith",
        answers=answers,
        version=1,
        created_at=datetime.now(timezone.utc),
        updated_at=datetime.now(timezone.utc),
    )

    result = docgen.render(org, policy)
    doc = docx.Document(io.BytesIO(result))

    metadata_table = doc.tables[1]
    assert metadata_table.rows[0].cells[1].text == "Render Test Org"
    assert metadata_table.rows[1].cells[1].text == "Jane Doe"
    assert metadata_table.rows[2].cells[1].text == "John Smith"

    body_text = "\n".join(p.text for p in doc.paragraphs)
    assert "Custom internal AI copilot" in body_text
    assert "Fulltime employees" in body_text

    # org name substituted inline, mid-sentence, in section 1.1 (not a
    # whole-paragraph placeholder like the checklist blocks are)
    assert 'Render Test Org (the "Organization") is committed' in body_text
    # default tier substituted inline in section 4.2, as its human label
    # ("Restricted"), not the raw stored slug ("restricted")
    assert "If uncertain about the tier — Restricted" in body_text

    # cover block: version/effective date/next review (+1 year) all landed
    cover_text = doc.tables[0].rows[0].cells[0].text
    assert "Version 1" in cover_text
    assert "Next Review:" in cover_text

    # governance owner subdoc table landed as a real table with the right content
    owner_tables = [t for t in doc.tables if t.rows[0].cells[0].text == "Name" and "Jane Doe" in t.rows[1].cells[0].text]
    assert len(owner_tables) == 1

    # optional dept-leads/legal-lead sections are left empty by default -
    # their whole subsection (heading included) must be omitted entirely
    assert "3.2" not in body_text
    assert "Department AI Leads" not in body_text
    assert "3.3" not in body_text
    assert "Legal & Compliance Lead" not in body_text

    # oversight row-loop table has all 8 default rows
    oversight_table = next(t for t in doc.tables if t.rows[0].cells[0].text == "Context")
    assert len(oversight_table.rows) == 9  # header + 8 rows

    # ampersand in a checklist label survives (autoescape regression check)
    tier_table = next(t for t in doc.tables if t.rows[0].cells[0].text == "Tier")
    assert "M&A activity" in tier_table.rows[2].cells[2].text


def test_optional_governance_sections_appear_when_filled() -> None:
    """3.2 (Department AI Leads) and 3.3 (Legal & Compliance Lead) are
    conditionally included - present with a real table when answered,
    completely absent (heading and all) when left blank. This covers the
    "present" half; the default-answers render above covers "absent"."""
    org = Org(id=uuid.uuid4(), name="Optional Sections Org")
    answers = default_answers()
    answers["q2_1_governance_owner"]["rows"] = [{"name": "Jane Doe", "title": "Chief AI Officer"}]
    answers["q2_4_approvers"]["rows"] = [{"name": "John Smith", "title": "CEO"}]
    answers["q3_9_default_tier"]["selected"] = "restricted"
    answers["q2_2_dept_leads"]["rows"] = [{"name": "Sam Lee", "department": "Engineering", "title": "Eng AI Lead"}]
    answers["q2_3_legal_lead"]["rows"] = [{"name": "Pat Wong", "department": "Legal", "title": "General Counsel"}]

    policy = Policy(
        id=uuid.uuid4(),
        org_id=org.id,
        status="draft",
        policy_owner_name="Jane Doe",
        approver_name="John Smith",
        answers=answers,
        version=1,
        created_at=datetime.now(timezone.utc),
        updated_at=datetime.now(timezone.utc),
    )

    result = docgen.render(org, policy)
    doc = docx.Document(io.BytesIO(result))
    body_text = "\n".join(p.text for p in doc.paragraphs)

    assert "3.2  Department AI Leads" in body_text
    assert "3.3  Legal & Compliance Lead" in body_text
    dept_tables = [t for t in doc.tables if t.rows[0].cells[0].text == "Name" and "Sam Lee" in t.rows[1].cells[0].text]
    legal_tables = [t for t in doc.tables if t.rows[0].cells[0].text == "Name" and "Pat Wong" in t.rows[1].cells[0].text]
    assert len(dept_tables) == 1
    assert len(legal_tables) == 1


def test_generate_rejects_incomplete_draft() -> None:
    org = _make_org("Policy Validation Test Org")
    try:
        with org_scoped_session(str(org.id)) as db:
            policy = Policy(org_id=org.id, answers=default_answers())
            db.add(policy)
            db.flush()
            db.refresh(policy)
            policy_id = str(policy.id)

        with pytest.raises(PolicyValidationError) as exc_info:
            service.generate(org, policy_id)

        assert "policy_owner_name" in exc_info.value.missing_required
        assert "approver_name" in exc_info.value.missing_required
        assert "q2_1_governance_owner" in exc_info.value.missing_required
        assert "q2_4_approvers" in exc_info.value.missing_required
        assert "q3_9_default_tier" in exc_info.value.missing_required
    finally:
        _delete_org(org)


def test_approve_promotes_draft_and_archives_previous_active() -> None:
    org = _make_org("Approval Test Org")
    try:
        approver = _make_user(org, "approver@example.com")
        first = _make_policy_with_document(org)
        second = _make_policy_with_document(org)

        approved_first = service.approve(org, str(first.id), approver)
        assert approved_first.status == "active"
        assert approved_first.version == 1
        assert approved_first.approved_by == approver.id

        approved_second = service.approve(org, str(second.id), approver)
        assert approved_second.status == "active"
        assert approved_second.version == 2

        refreshed_first = service.get_policy(org, str(first.id))
        assert refreshed_first is not None
        assert refreshed_first.status == "archived"
    finally:
        _delete_org(org)


def test_approve_caches_active_policy_as_markdown() -> None:
    org = _make_org("Approval Markdown Org")
    try:
        approver = _make_user(org, "md-approver@example.com")
        policy = _make_policy_with_document(org)

        approved = service.approve(org, str(policy.id), approver)
        assert approved.markdown_key is not None

        text = service.get_active_policy_text(org)
        assert text is not None
        assert "Test policy body text." in text
    finally:
        _delete_org(org)


def test_get_active_policy_text_backfills_missing_markdown_key() -> None:
    """A policy approved before markdown caching existed has no
    markdown_key yet - get_active_policy_text() should still work, by
    lazily extracting from the docx and backfilling the key for next
    time."""
    org = _make_org("Approval Backfill Org")
    try:
        approver = _make_user(org, "backfill-approver@example.com")
        policy = _make_policy_with_document(org)
        approved = service.approve(org, str(policy.id), approver)

        # Simulate a pre-Milestone-10 active policy: clear the cached key.
        with org_scoped_session(str(org.id)) as db:
            row = db.get(Policy, approved.id)
            row.markdown_key = None
            db.flush()

        text = service.get_active_policy_text(org)
        assert text is not None
        assert "Test policy body text." in text

        with org_scoped_session(str(org.id)) as db:
            row = db.get(Policy, approved.id)
            assert row.markdown_key is not None
    finally:
        _delete_org(org)


def test_has_active_policy() -> None:
    org = _make_org("Has Active Policy Org")
    try:
        assert service.has_active_policy(org) is False

        approver = _make_user(org, "hap-approver@example.com")
        policy = _make_policy_with_document(org)
        service.approve(org, str(policy.id), approver)

        assert service.has_active_policy(org) is True
    finally:
        _delete_org(org)


def test_approve_rejects_draft_with_no_document() -> None:
    org = _make_org("Approval Missing Doc Org")
    try:
        approver = _make_user(org, "approver2@example.com")
        with org_scoped_session(str(org.id)) as db:
            policy = Policy(org_id=org.id, answers={})
            db.add(policy)
            db.flush()
            db.refresh(policy)
            policy_id = str(policy.id)

        with pytest.raises(PolicyApprovalError):
            service.approve(org, policy_id, approver)
    finally:
        _delete_org(org)


def test_approve_rejects_a_policy_that_is_not_a_draft() -> None:
    org = _make_org("Approval Already Active Org")
    try:
        approver = _make_user(org, "approver3@example.com")
        policy = _make_policy_with_document(org)
        service.approve(org, str(policy.id), approver)

        with pytest.raises(PolicyApprovalError):
            service.approve(org, str(policy.id), approver)
    finally:
        _delete_org(org)


def _minimal_docx_bytes() -> bytes:
    buf = io.BytesIO()
    docx.Document().save(buf)
    return buf.getvalue()


def test_upload_policy_creates_a_draft_with_source_upload() -> None:
    org = _make_org("Upload Test Org")
    try:
        uploader = _make_user(org, "uploader@example.com")
        policy = service.upload_policy(org, uploader, "my-policy.docx", _minimal_docx_bytes())

        assert policy.status == "draft"
        assert policy.source == "upload"
        assert policy.storage_key is not None
        assert policy.generated_at is not None
    finally:
        _delete_org(org)


def test_upload_policy_rejects_a_docx_extension_with_non_docx_content() -> None:
    """The .docx-extension check alone is not enough - a plain text/binary
    blob renamed to end in .docx must still be rejected, not accepted and
    left to blow up later at approve() time."""
    org = _make_org("Upload Rejects Fake Content Org")
    try:
        uploader = _make_user(org, "uploader4@example.com")
        with pytest.raises(PolicyUploadError):
            service.upload_policy(org, uploader, "my-policy.docx", b"this is not a real docx file")
    finally:
        _delete_org(org)


def test_upload_policy_rejects_non_docx_filename() -> None:
    org = _make_org("Upload Rejects Filetype Org")
    try:
        uploader = _make_user(org, "uploader2@example.com")
        with pytest.raises(PolicyUploadError):
            service.upload_policy(org, uploader, "my-policy.pdf", b"not a docx")
    finally:
        _delete_org(org)


def test_upload_policy_rejects_oversized_file() -> None:
    org = _make_org("Upload Rejects Size Org")
    try:
        uploader = _make_user(org, "uploader3@example.com")
        with pytest.raises(PolicyUploadError):
            service.upload_policy(org, uploader, "my-policy.docx", b"0" * (MAX_UPLOAD_SIZE_BYTES + 1))
    finally:
        _delete_org(org)
